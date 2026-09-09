"""Build the PUBLIC-MINISTRY English writing-rubric fixture as a DOCX.

SOURCE: Israeli Ministry of Education, Module G (16582) / F (external 16584) WRITING RUBRIC,
as of Winter 2020. Fetched 2026-09-05 from
https://meyda.education.gov.il/files/Pop/0files/english/Chativa-Elyona/Bagrut/FGExternalInternal2020.pdf
Transcribed from the PDF's own text layer (PyMuPDF) into the shape a teacher's DOCX has: a
band table with one row per criterion, four band columns, and the points row beneath.
Nothing is invented; the wording is the ministry's.
"""
import sys
from docx import Document
from docx.shared import Pt

OUT = sys.argv[1]

CRITERIA = [
    ("CONTENT AND ORGANIZATION",
     ["fully on topic", "fully developed (main idea and supporting details)",
      "all elements of task are addressed", "information is relevant",
      "content is understood", "task is well organized and coherent"],
     ["partially on topic", "partially developed (main idea or supporting details)",
      "partially addresses elements of task", "information is partially relevant",
      "content is partially understood", "task is partially organized and coherent"],
     ["minimally on topic", "minimally developed (main idea or supporting details)",
      "minimally addresses elements of task", "information is minimally relevant",
      "content is minimally understood", "task is minimally organized and coherent"],
     ["not on topic", "not developed (main idea or supporting details)",
      "elements are not addressed *", "information is not relevant *",
      "content is not understood", "task is not organized and not coherent"],
     ["8", "5", "2", "0"]),
    ("VOCABULARY",
     ["correct use of varied and rich vocabulary",
      "appropriate use of instances of language chunks and phrases",
      "correct use of connecting words or phrases", "use of appropriate register"],
     ["correct use of basic, appropriate vocabulary",
      "partially appropriate use of instances of chunks and phrases.",
      "partial and correct use of connecting words or phrases",
      "occasional use of inappropriate register"],
     ["minimally correct use of basic, appropriate vocabulary",
      "minimally appropriate instances of chunks and phrases.",
      "minimal use of connecting words or phrases",
      "consistent use of inappropriate register"],
     ["incorrect use of words", "inappropriate use of chunks and phrases",
      "no use of connecting words or phrases", "consistent use of inappropriate register"],
     ["10", "6", "3", "0"]),
    ("LANGUAGE USE",
     ["correct use of basic tenses and/or language structures",
      "correct use of advanced language structures", "correct word order",
      "correct use of parts of speech, pronouns and prepositions"],
     ["correct use of basic tenses and/or language structures",
      "incorrect or no use of advanced language structures",
      "occasional instances of incorrect word order",
      "occasional incorrect use of parts of speech, pronouns & prepositions"],
     ["minimally correct use of basic tenses and/or language structures",
      "incorrect or no use of advanced language structures",
      "minimally correct word order",
      "minimally correct use of parts of speech, pronouns and prepositions"],
     ["Incorrect use of basic tenses and/or language structures",
      "incorrect or no use of advanced language structures", "incorrect word order",
      "incorrect use of parts of speech, pronouns and prepositions"],
     ["16", "10", "5", "0"]),
    ("MECHANICS",
     ["correct use of: spelling", "punctuation", "capitalization", "paragraphing",
      "no run-on sentences"],
     ["partially correct use of: spelling", "punctuation", "capitalization", "paragraphing",
      "some run-on sentences"],
     ["minimally correct use of: spelling", "punctuation", "capitalization", "paragraphing",
      "frequent run-on sentences"],
     ["Incorrect use of: spelling", "punctuation", "capitalization", "paragraphing",
      "consistent use of run-on sentences"],
     ["6", "4", "2", "0"]),
]

doc = Document()
doc.styles["Normal"].font.size = Pt(9)
doc.add_paragraph("MODULE G (16582) and F (external 16584) - WRITING RUBRIC - as of Winter 2020")

table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells,
                      ["CRITERIA", "CORRECT", "PARTIALLY CORRECT", "MINIMALLY CORRECT", "INCORRECT"]):
    cell.text = text

for name, correct, partial, minimal, incorrect, points in CRITERIA:
    row = table.add_row().cells
    row[0].text = name
    for cell, bullets in zip(row[1:], (correct, partial, minimal, incorrect)):
        cell.text = "\n".join(f"• {b}" for b in bullets)
    pts = table.add_row().cells
    pts[0].text = ""
    for cell, p in zip(pts[1:], points):
        cell.text = p

doc.add_paragraph("(Question = 40 points)")
doc.add_paragraph("GENERAL COMMENTS:")
for line in [
    "1. Markers can give in-between grades e.g. 7 pts.",
    "2. In cases when the topic of the reading passage (unseen) and writing task are similar "
    "and student copies from the text: If the writing task has been copied in its entirety from "
    "the reading passage - zero for the entire task.",
    "3. An entire composition will receive a zero when any of these criteria occur: there are "
    "fewer than 50 words.",
]:
    doc.add_paragraph(line)

doc.save(OUT)
print("wrote", OUT)
